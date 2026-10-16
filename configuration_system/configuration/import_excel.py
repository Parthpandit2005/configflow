import openpyxl
from pathlib import Path

from configuration.database import get_connection


# ============================================================
# FILE LOCATION
# ============================================================

BASE_DIR = Path(__file__).resolve().parent.parent

EXCEL_PATH = (
    BASE_DIR
    / "source_data"
    / "historical_configuration_knowledge.xlsx"
)


# ============================================================
# HELPER
# ============================================================

def clean(value):
    """
    Convert Excel cell values into clean strings.

    Empty cells become None.
    """

    if value is None:
        return None

    value = str(value).strip()

    if value == "":
        return None

    return value


# ============================================================
# IMPORT EXCEL
# ============================================================

def import_excel():

    print("=" * 70)
    print("IMPORTING CONFIGURATION KNOWLEDGE")
    print("=" * 70)

    if not EXCEL_PATH.exists():
        raise FileNotFoundError(
            f"Excel file not found:\n{EXCEL_PATH}"
        )

    print()
    print("Source file:")
    print(EXCEL_PATH)
    print()

    # --------------------------------------------------------
    # LOAD WORKBOOK
    # --------------------------------------------------------

    workbook = openpyxl.load_workbook(
        EXCEL_PATH,
        data_only=True
    )

    worksheet = workbook.active

    print("Worksheet:", worksheet.title)
    print("Rows     :", worksheet.max_row)
    print("Columns  :", worksheet.max_column)
    print()

    # --------------------------------------------------------
    # READ FIRST ROW
    # --------------------------------------------------------

    first_row = [
        clean(cell.value)
        for cell in worksheet[1]
    ]

    # --------------------------------------------------------
    # DETERMINE COLUMN STRUCTURE
    # --------------------------------------------------------

    headers = {}

    for index, header in enumerate(first_row):

        if header is not None:
            header = str(header).strip().lower()
            headers[header] = index

    # --------------------------------------------------------
    # NEW 5-COLUMN FORMAT
    # --------------------------------------------------------

    if "requirement" in headers:

        requirement_col = headers["requirement"]

        if "data type" in headers:
            data_type_col = headers["data type"]
        else:
            data_type_col = None

        if "access" in headers:
            access_col = headers["access"]
        else:
            raise ValueError(
                "Could not find 'access' column."
            )

        if "plc parameter name" in headers:
            plc_col = headers["plc parameter name"]
        else:
            raise ValueError(
                "Could not find 'PLC parameter name' column."
            )

        if "pc parameter name" in headers:
            pc_col = headers["pc parameter name"]

        elif "pc paramater name" in headers:
            pc_col = headers["pc paramater name"]

        else:
            raise ValueError(
                "Could not find 'PC parameter name' column."
            )

        start_row = 2

    # --------------------------------------------------------
    # FALLBACK FOR OLD 4-COLUMN FORMAT
    # --------------------------------------------------------

    else:

        print("No header row detected.")
        print("Assuming old 4-column format.")

        requirement_col = 0
        data_type_col = None
        access_col = 1
        plc_col = 2
        pc_col = 3

        start_row = 1

    # --------------------------------------------------------
    # DATABASE
    # --------------------------------------------------------

    connection = get_connection()
    cursor = connection.cursor()

    # --------------------------------------------------------
    # CURRENT REQUIREMENT
    #
    # Excel uses blank cells to mean:
    #
    # "Same requirement as above."
    # --------------------------------------------------------

    current_requirement = None

    # --------------------------------------------------------
    # STATISTICS
    # --------------------------------------------------------

    new_requirements = 0
    existing_requirements = 0

    new_parameters = 0
    existing_parameters = 0

    new_mappings = 0
    updated_mappings = 0
    existing_mappings = 0

    new_relationships = 0
    existing_relationships = 0

    # ========================================================
    # PROCESS ROWS
    # ========================================================

    for row_number in range(
        start_row,
        worksheet.max_row + 1
    ):

        row = [
            clean(cell.value)
            for cell in worksheet[row_number]
        ]

        # ----------------------------------------------------
        # SAFETY CHECK
        # ----------------------------------------------------

        if len(row) <= max(
            requirement_col,
            access_col,
            plc_col,
            pc_col
        ):
            continue

        requirement_value = row[requirement_col]

        data_type = (
            row[data_type_col]
            if data_type_col is not None
            and data_type_col < len(row)
            else None
        )

        access_type = row[access_col]
        plc_name = row[plc_col]
        pc_name = row[pc_col]

        # ----------------------------------------------------
        # DETERMINE REQUIREMENT
        # ----------------------------------------------------

        if requirement_value is not None:

            current_requirement = requirement_value

            if current_requirement.upper() == "DEFAULT":
                requirement_type = "DEFAULT"
            else:
                requirement_type = "SELECTABLE"

            # Check whether requirement already exists
            cursor.execute(
                """
                SELECT id, requirement_type
                FROM requirements
                WHERE name = ?
                """,
                (current_requirement,)
            )

            requirement = cursor.fetchone()

            if requirement is None:

                cursor.execute(
                    """
                    INSERT INTO requirements
                    (
                        name,
                        requirement_type
                    )
                    VALUES (?, ?)
                    """,
                    (
                        current_requirement,
                        requirement_type
                    )
                )

                new_requirements += 1

            else:

                existing_requirements += 1

        # ----------------------------------------------------
        # IGNORE ROWS BEFORE FIRST REQUIREMENT
        # ----------------------------------------------------

        if current_requirement is None:
            continue

        # ----------------------------------------------------
        # IGNORE EMPTY PARAMETER ROW
        # ----------------------------------------------------

        if plc_name is None and pc_name is None:
            continue

        # ----------------------------------------------------
        # PARAMETER NAME
        #
        # PLC name is the primary parameter name.
        #
        # If PLC name is missing, use PC name.
        # ----------------------------------------------------

        parameter_name = plc_name or pc_name

        # ----------------------------------------------------
        # CHECK PARAMETER
        # ----------------------------------------------------

        cursor.execute(
            """
            SELECT id, data_type
            FROM parameters
            WHERE name = ?
            """,
            (parameter_name,)
        )

        parameter = cursor.fetchone()

        if parameter is None:

            # ------------------------------------------------
            # NEW PARAMETER
            # ------------------------------------------------

            cursor.execute(
                """
                INSERT INTO parameters
                (
                    name,
                    data_type
                )
                VALUES (?, ?)
                """,
                (
                    parameter_name,
                    data_type
                )
            )

            parameter_id = cursor.lastrowid

            new_parameters += 1

        else:

            # ------------------------------------------------
            # EXISTING PARAMETER
            # ------------------------------------------------

            parameter_id = parameter[0]
            existing_parameters += 1

            # ------------------------------------------------
            # UPDATE DATA TYPE IF EXCEL PROVIDES ONE
            # ------------------------------------------------

            if data_type is not None:

                cursor.execute(
                    """
                    UPDATE parameters
                    SET data_type = ?
                    WHERE id = ?
                    """,
                    (
                        data_type,
                        parameter_id
                    )
                )

        # ====================================================
        # PARAMETER MAPPING
        # ====================================================

        cursor.execute(
            """
            SELECT
                id,
                access_type,
                plc_parameter_name,
                pc_parameter_name
            FROM parameter_mappings
            WHERE parameter_id = ?
              AND (
                    plc_parameter_name = ?
                    OR (
                        plc_parameter_name IS NULL
                        AND ? IS NULL
                    )
                  )
              AND (
                    pc_parameter_name = ?
                    OR (
                        pc_parameter_name IS NULL
                        AND ? IS NULL
                    )
                  )
            """,
            (
                parameter_id,
                plc_name,
                plc_name,
                pc_name,
                pc_name
            )
        )

        mapping = cursor.fetchone()

        if mapping is None:

            # ------------------------------------------------
            # NEW MAPPING
            # ------------------------------------------------

            cursor.execute(
                """
                INSERT INTO parameter_mappings
                (
                    parameter_id,
                    plc_parameter_name,
                    pc_parameter_name,
                    access_type
                )
                VALUES (?, ?, ?, ?)
                """,
                (
                    parameter_id,
                    plc_name,
                    pc_name,
                    access_type
                )
            )

            new_mappings += 1

        else:

            mapping_id = mapping[0]

            # ------------------------------------------------
            # CHECK WHETHER ACCESS TYPE CHANGED
            # ------------------------------------------------

            if mapping[1] != access_type:

                cursor.execute(
                    """
                    UPDATE parameter_mappings
                    SET access_type = ?
                    WHERE id = ?
                    """,
                    (
                        access_type,
                        mapping_id
                    )
                )

                updated_mappings += 1

            else:

                existing_mappings += 1

        # ====================================================
        # REQUIREMENT ID
        # ====================================================

        cursor.execute(
            """
            SELECT id
            FROM requirements
            WHERE name = ?
            """,
            (current_requirement,)
        )

        requirement = cursor.fetchone()

        if requirement is None:

            raise RuntimeError(
                f"Could not find requirement: "
                f"{current_requirement}"
            )

        requirement_id = requirement[0]

        # ====================================================
        # REQUIREMENT → PARAMETER
        # ====================================================

        cursor.execute(
            """
            SELECT 1
            FROM requirement_parameters
            WHERE requirement_id = ?
              AND parameter_id = ?
            """,
            (
                requirement_id,
                parameter_id
            )
        )

        relationship = cursor.fetchone()

        if relationship is None:

            cursor.execute(
                """
                INSERT INTO requirement_parameters
                (
                    requirement_id,
                    parameter_id
                )
                VALUES (?, ?)
                """,
                (
                    requirement_id,
                    parameter_id
                )
            )

            new_relationships += 1

        else:

            existing_relationships += 1

    # ========================================================
    # COMMIT
    # ========================================================

    connection.commit()

    # ========================================================
    # FINAL DATABASE COUNTS
    # ========================================================

    cursor.execute(
        """
        SELECT COUNT(*)
        FROM requirements
        """
    )

    total_requirements = cursor.fetchone()[0]

    cursor.execute(
        """
        SELECT COUNT(*)
        FROM parameters
        """
    )

    total_parameters = cursor.fetchone()[0]

    cursor.execute(
        """
        SELECT COUNT(*)
        FROM parameter_mappings
        """
    )

    total_mappings = cursor.fetchone()[0]

    cursor.execute(
        """
        SELECT COUNT(*)
        FROM requirement_parameters
        """
    )

    total_relationships = cursor.fetchone()[0]

    connection.close()

    # ========================================================
    # RESULT
    # ========================================================

    print()
    print("=" * 70)
    print("IMPORT COMPLETE")
    print("=" * 70)

    print()
    print("REQUIREMENTS")
    print("----------------------")
    print("New       :", new_requirements)
    print("Existing  :", existing_requirements)

    print()
    print("PARAMETERS")
    print("----------------------")
    print("New       :", new_parameters)
    print("Existing  :", existing_parameters)

    print()
    print("MAPPINGS")
    print("----------------------")
    print("New       :", new_mappings)
    print("Updated   :", updated_mappings)
    print("Existing  :", existing_mappings)

    print()
    print("RELATIONSHIPS")
    print("----------------------")
    print("New       :", new_relationships)
    print("Existing  :", existing_relationships)

    print()
    print("TOTAL DATABASE CONTENT")
    print("----------------------")
    print("Requirements :", total_requirements)
    print("Parameters   :", total_parameters)
    print("Mappings     :", total_mappings)
    print("Relationships:", total_relationships)

    print()
    print("Knowledge Base updated successfully.")


# ============================================================
# RUN
# ============================================================

if __name__ == "__main__":
    import_excel()